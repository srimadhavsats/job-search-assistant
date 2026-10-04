"""Link checks, watcher schedule and lock, alert selection."""
from datetime import datetime, timedelta

import httpx
import pytest

from jobhunter import notify, verify, watch
from jobhunter.lock import RunLock
from jobhunter.models import Job
from tests.conftest import mock_client_factory

BANDS = [{"min": 70, "label": "1 - Apply now"}, {"min": 55, "label": "2 - Good fit"}, {"min": 0, "label": "3 - Maybe"}]


# ---------------- verify ----------------
@pytest.mark.parametrize("html,result,flag", [
    ("<h1>Support Engineer</h1><p>No longer accepting applications</p>", "closed", ""),
    ("<h1>Support</h1><p>You must be authorized to work in the United States.</p>", "open", "Apply page requires work authorisation in the UNITED STATES"),
    ("<h1>Support</h1><p>Requirements. Based in the US.</p>", "open", "based in the us"),
    ("<h1>Support</h1><p>We hire globally, work from anywhere.</p>", "open", ""),
])
def test_page_verdict(sc, html, result, flag):
    r, f = verify.page_verdict(verify._visible_text(html), "Support Engineer", sc)
    assert r == result and flag.lower() in f.lower()


def test_apply_checks_pages_and_caps_locked_rows(monkeypatch, store, sc):
    pages = {"/open": "<p>Great crypto support job, remote</p>",
             "/us": "<p>Crypto support. Must be authorized to work in the US.</p>"}

    def handler(req):
        if req.url.path == "/gone":
            return httpx.Response(404)
        return httpx.Response(200, text=pages.get(req.url.path, ""))

    monkeypatch.setattr(verify, "http_client", mock_client_factory(handler))
    rows = [{"Key": k, "Title": "Support Engineer", "Score": 80, "Status": "New", "Source": "web3.career",
             "Apply link": f"https://board.example{path}", "Flags": ""} for k, path in
            (("k1", "/open"), ("k2", "/us"), ("k3", "/gone"))]
    rows.append({"Key": "k4", "Title": "X", "Score": 80, "Status": "New", "Source": "Company careers: Binance",
                 "Apply link": "https://jobs.lever.co/binance/1"})
    stats = verify.apply(rows, store, sc, BANDS, log=lambda m: None, delay=0)
    by = {r["Key"]: r for r in rows}
    assert stats["checked"] == 3                              # company-career rows need no visit
    assert by["k1"]["Link check"].startswith("Open") and by["k1"]["Score"] == 80
    assert by["k2"]["Score"] <= sc["region_locked_max_score"] and "Apply page requires" in by["k2"]["Flags"]
    assert by["k3"]["Link check"].startswith("Closed") and by["k3"]["Score"] <= 30
    # a second pass the same day opens nothing again
    assert verify.apply(rows, store, sc, BANDS, log=lambda m: None, delay=0)["checked"] == 0


# ---------------- watcher ----------------
def test_due_sources_and_backoff(store):
    cfgs = [{"name": "Fast", "type": "rss", "every_minutes": 20}, {"name": "Slow", "type": "naukri", "every_minutes": 240}]
    now = datetime(2026, 9, 25, 12, 0)
    assert [c["name"] for c in watch.due_sources(cfgs, store, now)] == ["Fast", "Slow"]   # never ran
    store.record_source("Fast", True, "ok", 10, 20, now)
    store.record_source("Slow", True, "ok", 10, 240, now)
    assert watch.due_sources(cfgs, store, now + timedelta(minutes=10)) == []
    assert [c["name"] for c in watch.due_sources(cfgs, store, now + timedelta(minutes=21))] == ["Fast"]
    # failures retry after 15, 30, 60 minutes… capped at 6 hours
    waits = []
    for i in range(7):
        store.record_source("Fast", False, "FAILED", 0, 20, now)
        waits.append((datetime.fromisoformat(store.source_state("Fast")["next_due"]) - now).total_seconds() / 60)
    assert waits[:4] == [15, 30, 60, 120] and waits[-1] == 360
    store.record_source("Fast", True, "ok", 5, 20, now)
    assert store.source_state("Fast")["fails"] == 0


def test_lock_is_exclusive(tmp_path):
    a, b = RunLock(tmp_path / "run.lock"), RunLock(tmp_path / "run.lock")
    assert a.acquire()
    assert not b.acquire(wait_seconds=0)
    a.release()
    assert b.acquire()
    b.release()


# ---------------- alerts ----------------
def test_worth_alert_rules():
    cfg = {"min_score": 60, "freelance_min_score": 65, "remote_bonus": 5, "only_applicable": True}
    base = {"Status": "New", "Score": 62, "Location": "Bengaluru, India", "Flags": ""}
    assert notify.worth_alert(base, cfg)
    assert not notify.worth_alert({**base, "Score": 58}, cfg)
    assert notify.worth_alert({**base, "Score": 57, "Location": "Remote (Worldwide)", "Work mode": "Remote"}, cfg)
    assert not notify.worth_alert({**base, "Flags": "Remote but US-only — you likely can't apply from India"}, cfg)
    assert not notify.worth_alert({**base, "Status": "Applied"}, cfg)
    assert not notify.worth_alert({**base, "_kind": "freelance", "Score": 62}, cfg)


def test_new_jobs_alerts_once(monkeypatch, store):
    shown = []
    monkeypatch.setattr(notify, "windows_popup", lambda *a, **k: shown.append(a) or True)
    cfg = {"min_score": 60, "windows_popup": True, "max_popups_per_check": 4, "telegram": {}}
    rows = {"k1": {"Key": "k1", "Status": "New", "Score": 80, "Title": "KYC Analyst: EDD", "Company": "Binance",
                   "Location": "Remote", "Apply link": "https://x/1&y=2", "Flags": ""}}
    assert notify.new_jobs(rows, {"k1"}, store, log=lambda m: None, cfg=cfg) == 1
    assert notify.new_jobs(rows, {"k1"}, store, log=lambda m: None, cfg=cfg) == 0
    title, body = shown[0][0], shown[0][1]
    assert ":" not in title and ":" not in body


def test_plain_text_has_no_banned_marks():
    assert notify.plain("Compliance: KYC; AML — Lead | Remote") == "Compliance, KYC, AML, Lead, Remote"


def test_ats_links_are_checked_through_their_apis(monkeypatch, store, sc):
    """VC boards link to Ashby/Lever/Greenhouse pages without a description (Wormhole, 2026-09-25)."""
    ashby = {"jobs": [{"id": "11111111-2222-3333-4444-555555555555", "location": "New York",
                       "descriptionPlain": "Trading operations for a crypto bridge."}]}

    def handler(req):
        u = str(req.url)
        if "api.ashbyhq.com" in u:
            return httpx.Response(200, json=ashby)
        if "api.lever.co" in u:
            return httpx.Response(404)
        if "boards-api.greenhouse.io" in u:
            return httpx.Response(200, json={"location": {"name": "Remote"}, "content": "&lt;p&gt;You must be authorized to work in the UK&lt;/p&gt;"})
        return httpx.Response(500)

    monkeypatch.setattr(verify, "http_client", mock_client_factory(handler))
    rows = [
        {"Key": "a", "Title": "Trading Operations Associate", "Score": 75, "Status": "New", "Source": "VC job boards: Solana",
         "Apply link": "https://jobs.ashbyhq.com/wormholelabs/11111111-2222-3333-4444-555555555555"},
        {"Key": "b", "Title": "Ops", "Score": 75, "Status": "New", "Source": "VC job boards: Dragonfly",
         "Apply link": "https://jobs.lever.co/acme/aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"},
        {"Key": "c", "Title": "Support", "Score": 75, "Status": "New", "Source": "JobStash",
         "Apply link": "https://job-boards.greenhouse.io/acme/jobs/123"},
        {"Key": "d", "Title": "Support", "Score": 75, "Status": "New", "Source": "VC job boards: Solana",
         "Apply link": "https://jobs.ashbyhq.com/wormholelabs/99999999-2222-3333-4444-555555555555"},
    ]
    verify.apply(rows, store, sc, BANDS, log=lambda m: None, delay=0)
    by = {r["Key"]: r for r in rows}
    assert "New York" in by["a"]["Flags"] and by["a"]["Score"] <= 45
    assert by["b"]["Link check"].startswith("Closed")
    assert "work authorisation in the UK" in by["c"]["Flags"]
    assert by["d"]["Link check"].startswith("Closed")


# ---- 2026-09-26: laptop off at 10:00, Wi-Fi drops, false "probably blocked" alarms ----
def test_morning_pass_due():
    at = lambda h, m: datetime(2026, 9, 26, h, m)
    assert not watch.morning_pass_due(at(8, 0), "")                       # too early
    assert watch.morning_pass_due(at(11, 15), "2026-09-25T20:28:00")        # laptop off at 10:00, first check after boot
    assert not watch.morning_pass_due(at(11, 45), "2026-09-26T11:15:00")    # already done today
    assert watch.morning_pass_due(at(9, 30), "")


def test_catch_up_only_after_a_long_gap():
    now = datetime(2026, 9, 26, 11, 15)
    js = {"name": "JobStash", "every_minutes": 60, "watch": {"publication_date": "today"}}
    plain = {"name": "Company careers", "every_minutes": 30}
    assert watch.needs_catch_up(js, {"last_ok": "2026-09-26T04:05:00"}, now)       # 7 h offline
    assert not watch.needs_catch_up(js, {"last_ok": "2026-09-26T10:30:00"}, now)
    assert not watch.needs_catch_up(plain, {"last_ok": "2026-09-25T10:00:00"}, now)  # no quick settings to widen


def test_quick_checks_have_their_own_history():
    from jobhunter import pipeline
    assert pipeline.usual_key({"name": "LinkedIn", "_quick": True}) == "LinkedIn [quick]"
    assert pipeline.usual_key({"name": "LinkedIn"}) == "LinkedIn"


def test_postpone_does_not_count_a_failure(store):
    now = datetime(2026, 9, 26, 4, 40)
    store.record_source("Naukri", False, "FAILED", 0, 240, now)
    store.postpone("Naukri", 15, "offline during check", now)
    st = store.source_state("Naukri")
    assert st["fails"] == 1 and st["next_due"] == "2026-09-26T04:55:00"   # not 2, and back in 15 min


def test_lock_is_free_after_the_holder_is_killed(tmp_path):
    """Power cut or crash mid-check: the OS drops the lock, so the next check is never blocked."""
    import subprocess
    import sys
    import time as real_time
    lock_path = tmp_path / "run.lock"
    code = ("import sys, time; sys.path.insert(0, r'%s'); from jobhunter.lock import RunLock; "
            "l = RunLock(r'%s'); assert l.acquire(); print('held', flush=True); time.sleep(60)") % (
        str(__import__('tests.conftest', fromlist=['ROOT']).ROOT), str(lock_path))
    proc = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.PIPE, text=True)
    assert proc.stdout.readline().strip() == "held"
    assert not RunLock(lock_path).acquire(wait_seconds=0)      # really held by the other process
    proc.kill()
    proc.wait(timeout=10)
    again = RunLock(lock_path)
    assert again.acquire(wait_seconds=0)
    again.release()


def test_interrupted_save_leaves_the_workbook_untouched(tmp_path, monkeypatch):
    """A crash while writing only ever hits the temp file, which the next save replaces."""
    from openpyxl import Workbook, load_workbook
    from jobhunter import excel
    target = tmp_path / "Jobs.xlsx"
    wb = Workbook()
    wb.active["A1"] = "your data"
    wb.save(target)
    (tmp_path / "Jobs.saving.tmp").write_bytes(b"half written")       # left by a power cut
    new = Workbook()
    new.active["A1"] = "next save"

    def boom(self, filename):
        raise OSError("power cut")
    monkeypatch.setattr(Workbook, "save", boom)
    with pytest.raises(OSError):
        excel._save_safely(new, target)
    assert load_workbook(target).active["A1"].value == "your data"
    monkeypatch.undo()
    excel._save_safely(new, target)
    assert load_workbook(target).active["A1"].value == "next save"
    assert not (tmp_path / "Jobs.saving.tmp").exists()


def test_vc_portfolio_page_text_is_saved_and_non_crypto_rows_are_capped(monkeypatch, store, sc):
    """Veem (payments, on Pantera's board) was #4 in the Apply Queue on 2026-09-26 with no crypto word anywhere."""
    from jobhunter.scoring import PORTFOLIO_NOT_CRYPTO
    filler = " Full time role, apply with your CV and a short note about your experience." * 3
    pages = {"/veem": "<p>Veem transforms cross-border payments. REST APIs, OAuth, webhooks, Postman.</p>" + filler,
             "/wallet": "<p>Support our crypto wallet users with stuck blockchain transactions.</p>" + filler}
    monkeypatch.setattr(verify, "http_client", mock_client_factory(
        lambda req: httpx.Response(200, text=pages.get(req.url.path, ""))))
    rows = [{"Key": k, "Title": "Technical Support Specialist", "Company": "Acme", "Score": 80, "Status": "New",
             "Source": src, "Apply link": f"https://careers.example{path}", "Flags": "", "Priority": "1 - Apply now"}
            for k, src, path in (("veem", "VC job boards: Pantera", "/veem"), ("wallet", "VC job boards: Solana", "/wallet"),
                                 ("board", "web3.career", "/veem"))]
    verify.apply(rows, store, sc, BANDS, log=lambda m: None, delay=0, hint_sources=("VC job boards",))
    by = {r["Key"]: r for r in rows}
    assert "cross-border payments" in store.text("veem")          # saved, so scoring uses it on the next fetch
    assert by["veem"]["Score"] <= sc["weak_web3_max_score"] and PORTFOLIO_NOT_CRYPTO in by["veem"]["Flags"]
    assert by["wallet"]["Score"] == 80 and not by["wallet"]["Flags"]
    assert by["board"]["Score"] == 80 and store.text("board") == ""   # other sources are left alone


def test_saved_page_text_fills_vc_listings_before_scoring(store):
    from jobhunter.pipeline import fill_saved_texts
    hint = Job(source="VC job boards: Pantera", title="Support", url="https://x/1", company="Veem", web3_hint=True)
    other = Job(source="LinkedIn", title="Support", url="https://x/2", company="Other")
    store.save_texts([(hint.key, "cross-border payments " * 20), (other.key, "should not be used " * 20)])
    assert fill_saved_texts([({}, [hint, other], "ok")], store) == 1
    assert hint.description.startswith("cross-border") and other.description == ""
