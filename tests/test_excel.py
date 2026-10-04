"""Workbook: tabs, status sync between tabs, side files, CV links (jobhunter/excel.py)."""
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from jobhunter import excel
from jobhunter.models import Job

TODAY = date(2026, 9, 25)
BANDS = [{"min": 70, "label": "1 - Apply now"}, {"min": 55, "label": "2 - Good fit"}, {"min": 0, "label": "3 - Maybe"}]


def job(title, company="Acme", score=75, location="Remote", work_mode="Remote", flags="", kind="job", source="web3.career"):
    j = Job(source=source, title=title, url=f"https://example.org/{title.replace(' ', '-')}", company=company,
            location=location, work_mode=work_mode, kind=kind)
    j.score, j.flags = score, flags
    j.priority = next(b["label"] for b in BANDS if score >= b["min"])
    return j


@pytest.fixture(autouse=True)
def not_open(monkeypatch):
    monkeypatch.setattr(excel, "_is_open", lambda p: False)


def save(path, jobs, today=TODAY, **kw):
    fs = {j.key: kw.pop("first_seen", today.isoformat()) for j in jobs}
    return excel.save(path, jobs, fs, today, [[f"{today} 10:00", "test", len(jobs), len(jobs), 0, "ok"]], BANDS, **kw)


def rows(path, sheet):
    wb = load_workbook(path)
    return excel._read(wb[sheet]) if sheet in wb.sheetnames else {}


def set_cell(path, sheet, key, column, value):
    wb = load_workbook(path)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        if r[header.index("Key")].value == key:
            r[header.index(column)].value = value
    wb.save(path)


def test_tabs_and_order(tmp_path):
    p = tmp_path / "w.xlsx"
    save(p, [job("Support Engineer"), job("Gig", kind="freelance", score=70)])
    names = load_workbook(p).sheetnames
    assert names[:7] == ["Today", "Apply Queue", "Follow Ups", "Remote Jobs", "Freelance", "All Jobs", "2026-09-25"]
    assert names[-1] == "Run Log"


def test_status_edited_in_any_tab_reaches_all_jobs(tmp_path):
    p = tmp_path / "w.xlsx"
    a, b, c = job("Support Engineer"), job("KYC Analyst", score=80), job("Ops Analyst", score=72)
    save(p, [a, b, c])
    set_cell(p, "Remote Jobs", a.key, "Status", "Applied")
    set_cell(p, "Apply Queue", b.key, "Status", "Skip")
    set_cell(p, "2026-09-25", c.key, "Notes", "asked Priya for a referral")
    save(p, [a, b, c])
    all_jobs = rows(p, "All Jobs")
    assert all_jobs[a.key]["Status"] == "Applied"
    assert all_jobs[b.key]["Status"] == "Skip"
    assert all_jobs[c.key]["Notes"] == "asked Priya for a referral"
    # every tab shows the same values afterwards
    assert rows(p, "2026-09-25")[a.key]["Status"] == "Applied"
    assert a.key not in rows(p, "Apply Queue") and b.key not in rows(p, "Apply Queue")


def test_applied_in_all_jobs_is_not_undone_by_day_tab(tmp_path):
    """Before 2026-09-25 a 'New' on the same day's tab overwrote Applied set in All Jobs."""
    p = tmp_path / "w.xlsx"
    a = job("Support Engineer")
    save(p, [a])
    set_cell(p, "All Jobs", a.key, "Status", "Applied")
    save(p, [a])
    assert rows(p, "All Jobs")[a.key]["Status"] == "Applied"
    assert rows(p, "2026-09-25")[a.key]["Status"] == "Applied"


def test_change_status_again_later(tmp_path):
    p = tmp_path / "w.xlsx"
    a = job("Support Engineer")
    save(p, [a])
    set_cell(p, "All Jobs", a.key, "Status", "Applied")
    save(p, [a])
    set_cell(p, "All Jobs", a.key, "Status", "Interview")      # other tabs still say Applied
    save(p, [a], today=TODAY + timedelta(days=1))
    assert rows(p, "All Jobs")[a.key]["Status"] == "Interview"
    assert rows(p, "2026-09-25")[a.key]["Status"] == "Interview"


def test_status_on_older_day_tab_survives_midnight(tmp_path):
    p = tmp_path / "w.xlsx"
    a = job("Support Engineer")
    save(p, [a])
    set_cell(p, "2026-09-25", a.key, "Status", "Applied")
    save(p, [a], today=TODAY + timedelta(days=1), first_seen=TODAY.isoformat())
    assert rows(p, "All Jobs")[a.key]["Status"] == "Applied"


def test_remote_tab_only_remote_and_open_to_india(tmp_path):
    p = tmp_path / "w.xlsx"
    ok = job("Support Engineer", location="Remote (Worldwide)")
    locked = job("Ops Analyst", location="Remote - USA", flags="Remote but USA-only — you likely can't apply from India", score=45)
    onsite = job("KYC Analyst", location="Bengaluru, India", work_mode="")
    save(p, [ok, locked, onsite])
    remote = rows(p, "Remote Jobs")
    assert ok.key in remote and locked.key not in remote and onsite.key not in remote


def test_apply_queue_has_cv_referrer_and_message(tmp_path):
    p = tmp_path / "w.xlsx"
    a = job("KYC Analyst", company="Binance", score=80)

    def enrich(rs):
        for r in rs:
            r["CV"], r["_cv_link"] = "TM and AML CV (custom)", "cvs/CV_Binance_KYC_Analyst.pdf"
    save(p, [a, job("Low", score=40)], enrich=enrich)
    wb = load_workbook(p)
    ws = wb["Apply Queue"]
    header = [c.value for c in ws[1]]
    row = [c for c in ws[2]]
    assert ws.max_row == 2
    assert row[header.index("CV")].hyperlink.target == "cvs/CV_Binance_KYC_Analyst.pdf"
    assert "linkedin.com/search/results/people" in row[header.index("Find referrer")].hyperlink.target
    assert "Binance" in row[header.index("Referral message")].value
    for col in ("Key", "Saved status", "Saved notes"):
        assert ws.column_dimensions[ws.cell(1, header.index(col) + 1).column_letter].hidden


def test_side_file_rows_and_edits_are_merged(tmp_path):
    p = tmp_path / "Jobs.xlsx"
    a, b = job("Support Engineer"), job("New While Open", score=71)
    save(p, [a])
    set_cell(p, "All Jobs", a.key, "Status", "Applied")         # edited in the main workbook while it was open
    side = tmp_path / "Jobs_2026-09-25_1900.xlsx"
    import shutil
    shutil.copy(p, side)
    set_cell(side, "All Jobs", a.key, "Status", "New")           # the side copy was built before the edit
    set_cell(side, "All Jobs", a.key, "Saved status", "New")
    wb = load_workbook(side)                                      # side file also has a job found meanwhile
    save(side, [b])
    save(p, [])
    all_jobs = rows(p, "All Jobs")
    assert all_jobs[a.key]["Status"] == "Applied" and b.key in all_jobs
    assert not side.exists()


def test_run_log_keeps_14_days(tmp_path):
    p = tmp_path / "w.xlsx"
    old = TODAY - timedelta(days=20)
    excel.save(p, [], {}, old, [[f"{old} 10:00", "x", 1, 1, 0, "ok"]], BANDS)
    excel.save(p, [], {}, TODAY, [[f"{TODAY} 10:00", "x", 1, 1, 0, "ok"]], BANDS)
    ws = load_workbook(p)["Run Log"]
    assert ws.max_row == 2 and str(ws.cell(2, 1).value).startswith(str(TODAY))


def test_not_seen_since_needs_a_healthy_source(tmp_path):
    p = tmp_path / "w.xlsx"
    a = job("Support Engineer", source="Naukri")
    save(p, [a], today=TODAY - timedelta(days=5))
    # Naukri has been blocked since: no "maybe closed" flag
    save(p, [], source_last_ok={"Naukri": (TODAY - timedelta(days=5)).isoformat()})
    assert "not seen since" not in str(rows(p, "All Jobs")[a.key].get("Flags") or "")
    # Naukri ran fine today and did not list it: flagged
    save(p, [], source_last_ok={"Naukri": TODAY.isoformat()})
    assert "not seen since" in str(rows(p, "All Jobs")[a.key].get("Flags") or "")


def test_resolve_prefers_edits_over_untouched():
    master = {"Status": "Applied", "Saved status": "Applied", "Notes": "", "Saved notes": "", "_has_saved": True}
    edited = {"Status": "Interview", "Saved status": "Applied", "Notes": "", "Saved notes": "", "_has_saved": True}
    untouched = {"Status": "Applied", "Saved status": "Applied", "_has_saved": True}
    assert excel.resolve_user_fields(master, [untouched, edited])[0] == "Interview"
    legacy_master = {"Status": "New"}
    legacy_day = {"Status": "Applied"}
    assert excel.resolve_user_fields(legacy_master, [legacy_day])[0] == "Applied"


def test_status_date_and_follow_ups(tmp_path):
    p = tmp_path / "w.xlsx"
    a, b = job("KYC Analyst", company="Binance"), job("Support Engineer")
    save(p, [a, b], today=TODAY - timedelta(days=8))
    set_cell(p, "Apply Queue", a.key, "Status", "Applied")
    save(p, [a, b], today=TODAY - timedelta(days=7), first_seen=(TODAY - timedelta(days=8)).isoformat())
    assert rows(p, "All Jobs")[a.key]["Status date"] == (TODAY - timedelta(days=7)).isoformat()
    save(p, [a, b], today=TODAY, first_seen=(TODAY - timedelta(days=8)).isoformat())
    fu = rows(p, "Follow Ups")
    assert a.key in fu and b.key not in fu
    assert fu[a.key]["Days since"] == 7 and "KYC Analyst at Binance" in fu[a.key]["Follow-up message"]
    set_cell(p, "Follow Ups", a.key, "Status", "Interview")
    save(p, [a, b], today=TODAY, first_seen=(TODAY - timedelta(days=8)).isoformat())
    assert rows(p, "All Jobs")[a.key]["Status"] == "Interview" and a.key not in rows(p, "Follow Ups")


def test_queue_is_jobs_only_and_skips_scams_and_applied_duplicates(tmp_path):
    p = tmp_path / "w.xlsx"
    applied = job("Quality Analyst - Blockchain Applications", company="JP Morgan Services India Pvt Ltd", score=72)
    dup = job("Quality Analyst - Blockchain Applications", company="JPMorgan Chase Bank", score=72, source="Naukri")
    maybe = job("Quality Analyst", company="JPMorganChase", score=72, source="Indeed")
    scam = job("Crypto Support Agent", score=80, flags="⚠ SCAM? (registration fee)")
    gig = job("Research Thread", kind="freelance", score=85)
    save(p, [applied])
    set_cell(p, "All Jobs", applied.key, "Status", "Applied")
    save(p, [applied, dup, maybe, scam, gig])
    q = rows(p, "Apply Queue")
    assert dup.key not in q and scam.key not in q and gig.key not in q
    assert maybe.key in q and "Maybe the same job" in q[maybe.key]["Flags"]


def test_queue_skips_senior_below_70_and_old_posts():
    from jobhunter.excel import queue_ok
    base = {"Status": "New", "Score": 62, "Flags": "", "Why it matched": "title B: analyst"}
    assert queue_ok(base)
    assert not queue_ok({**base, "Why it matched": "title B: analyst | senior title"})
    assert queue_ok({**base, "Score": 75, "Why it matched": "title A: aml | senior title"})
    assert not queue_ok({**base, "Flags": "Old post (2024) — may be stale/evergreen"})
